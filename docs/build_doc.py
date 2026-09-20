# Builds the Reality CRM Feature Guide (.docx) with the detailed feature list
# and the representative UI mockups.
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = r"D:\Data\Projects\Reality CRM"
SHOTS = os.path.join(ROOT, "docs", "screenshots")
OUT = os.path.join(ROOT, "Reality-CRM-Feature-Guide.docx")

BRAND = RGBColor(0x4F, 0x46, 0xE5)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUT = RGBColor(0x64, 0x74, 0x8B)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = INK

def h1(t):
    p = doc.add_heading(level=1); r = p.add_run(t); r.font.color.rgb = BRAND; r.font.size = Pt(17); return p
def h2(t):
    p = doc.add_heading(level=2); r = p.add_run(t); r.font.color.rgb = INK; r.font.size = Pt(13); return p
def para(t, size=10.5, color=INK, italic=False, bold=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.font.size = Pt(size); r.font.color.rgb = color
    r.italic = italic; r.bold = bold; return p
def bullet(label, detail):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(label + " — "); r.bold = True; r.font.size = Pt(10.5)
    r2 = p.add_run(detail); r2.font.size = Pt(10.5); return p
def caption(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = MUT
def img(name, cap):
    path = os.path.join(SHOTS, name)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(6.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption(cap)

# ── Title ──
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Reality CRM"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = BRAND
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Feature Guide & Product Overview"); r.font.size = Pt(14); r.font.color.rgb = MUT
m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = m.add_run("A custom Real-Estate CRM for mid-to-large developers\nCampaign → Lead → Account → Opportunity → Site Visit → Quotation → Booking → Demand & Receipt")
r.font.size = Pt(10.5); r.font.color.rgb = MUT
doc.add_paragraph()

# ── Executive summary ──
h1("1. Executive Summary")
para("Reality CRM is a purpose-built, non-Salesforce CRM for residential real-estate developers. It manages the entire "
     "sales lifecycle — from marketing campaigns and lead capture through site visits, quotations, bookings, and post-sale "
     "demand & receipt collection — on top of a multi-tower inventory engine. The platform is multi-project from day one, "
     "RERA-aware, and built around strict role-based access, row-level data scoping, and an append-only audit trail.")
para("It is delivered as two applications: an Express/Prisma REST API and a React (Vite) single-page app, backed by "
     "PostgreSQL and Redis. A Channel-Partner (broker) portal integrates with the CRM via a secure server-to-server intake API.")

# ── At a glance ──
h1("2. Platform at a Glance")
rows = [
    ("Sales lifecycle", "Campaigns, Leads, Accounts, Opportunities, Site Visits, Quotations, Bookings, Demand & Receipts"),
    ("Inventory", "Projects → Towers → Units with a colour-coded availability grid, holds, and price revisions"),
    ("Access control", "6 roles, dot-namespaced permissions, row-level scoping, append-only audit log"),
    ("Money", "Stored as BigInt paise (never floats); rupee input converted at the edge; lossless on the client"),
    ("Multi-project & RERA", "Every record carries a project context; RERA fields on Project and Booking"),
    ("Extensibility", "Admin-defined custom fields, 6 UI themes, and a broker-portal integration"),
]
tbl = doc.add_table(rows=0, cols=2); tbl.style = "Light Grid Accent 1"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in rows:
    c = tbl.add_row().cells
    rr = c[0].paragraphs[0].add_run(k); rr.bold = True; rr.font.size = Pt(10)
    c[1].paragraphs[0].add_run(v).font.size = Pt(10)

# ── Screens ──
h1("3. Key Screens")
para("The following are representative interface mockups that illustrate the layout and visual design of the application. "
     "They are illustrative renders, not live screen captures.", size=9.5, color=MUT, italic=True)
img("01-login.png", "Figure 1 — Sign-in screen with role-based demo accounts and the branded “site plan” panel.")
img("02-dashboard.png", "Figure 2 — Dashboard: KPI cards, recent lead activity, pipeline value, and upcoming site visits.")
img("03-leads-kanban.png", "Figure 3 — Leads: drag-and-drop Kanban pipeline with AI lead scores and stage columns.")

# ── Detailed feature list ──
h1("4. Detailed Feature List")

h2("4.1 Authentication, Roles & Security")
bullet("JWT authentication", "Access/refresh token flow; the SPA silently refreshes on 401 and retries the original request.")
bullet("Role-based access control (RBAC)", "Six roles — Super Admin, CRM Admin, Sales Agent, Channel Partner, Project Manager, Finance — mapped to dot-namespaced permissions (e.g. leads.read). Super Admin bypasses all checks.")
bullet("Row-level scoping", "Admins/PM/Finance see everything; agents see only the records they own; channel partners are scoped to their partner account.")
bullet("Append-only audit log", "Every create/update/delete is written to an immutable AuditLog (never updated or deleted) with old/new values, user, module, and entity.")
bullet("Admin console", "User management (create, deactivate, reset password), role-to-permission editor, and a searchable, paginated audit log.")

h2("4.2 Campaigns & Lead Capture")
bullet("Campaigns", "Track marketing campaigns by type and status with per-source attribution; every lead can trace back to a campaign source.")
bullet("Lead management", "Capture name, mobile, email, source channel, interest type, BHK, budget range, and preferred location; automatic round-robin owner assignment for admin-created leads.")
bullet("Bulk import", "Paste-to-import leads (name, mobile, email) for quick onboarding of lists.")
bullet("Drag-and-drop Kanban", "Move leads across stages (New → Contacted → Qualified → Site Visit → Negotiation) with optimistic updates and instant feedback; a list view is also available.")
bullet("Activity timeline", "Log calls, emails, WhatsApp, and notes; every stage change is recorded as a timeline event.")
bullet("Lead conversion", "Converting a lead spawns an Account + Contact + Opportunity in a single transaction.")

h2("4.3 Accounts & Contacts")
bullet("Account 360", "Companies, individuals, partners, and investors with PAN/GSTIN, address, and a tabbed view of contacts, leads, opportunities, bookings, and documents.")
bullet("KYC workflow", "KYC status transitions (Pending → Submitted → Verified / Rejected); a verified KYC is required before a booking can be confirmed.")
bullet("Channel-partner accounts", "Empanelment status and commission percentage for broker accounts.")

h2("4.4 Opportunities")
bullet("Pipeline & forecast", "Kanban pipeline by stage with deal value, probability, and a forecast/weighted view.")
bullet("One open opportunity per unit", "Soft block with manager override to prevent double-selling a unit.")
bullet("Won locks the unit", "Confirming a booking flips the linked opportunity to WON.")

h2("4.5 Site Visits")
bullet("Scheduling", "Schedule visits against a lead and project; visits cannot be scheduled in the past; walk-ins supported.")
bullet("Status & feedback", "Status flow (Scheduled → Confirmed → Completed / No-show / Cancelled); feedback is required before a visit can be marked Completed.")
bullet("Calendar view", "A calendar of upcoming and past visits.")

h2("4.6 Inventory Engine")
bullet("Multi-tower model", "Projects → Towers → Units, each unit carrying type, status, and pricing.")
bullet("Colour-coded grid", "A floor × unit grid that visualises availability (available, blocked, booked, registered).")
bullet("Time-limited holds", "Only AVAILABLE units can be held; holds expire automatically.")
bullet("Price revisions", "Versioned price changes that skip BOOKED/REGISTERED units.")

h2("4.7 Quotations")
bullet("Itemised quotes", "Line-item pricing with GST per line; generated, branded, print-ready PDFs.")
bullet("Versioning & lock", "Accepting a quotation locks its pricing while retaining prior versions.")

h2("4.8 Bookings")
bullet("Multi-step wizard", "Guided booking flow that requires a verified-KYC account and an available unit.")
bullet("Sequential numbering", "Atomic counter generates gap-free booking/receipt/quote/demand numbers.")
bullet("Dual-approval cancellation", "Cancelling a booking needs both CRM-Admin and Finance approval.")
bullet("Allotment letter", "Branded, print-ready allotment letter PDF.")

h2("4.9 Finance — Demand & Receipts")
bullet("Payment plans", "Construction-linked / time-linked / custom plans whose milestones must sum to 100%.")
bullet("Demand schedules", "GST-itemised demand letters; overdue demands escalate automatically.")
bullet("Receipts & reconciliation", "Receipts cannot exceed outstanding (with an advance override) and auto-reconcile oldest-first; branded receipt PDFs.")

h2("4.10 Reports & Dashboard")
bullet("KPI dashboard", "At-a-glance leads, site visits, bookings, collections, and pipeline value (fan-out KPI queries).")
bullet("Saved reports", "Build and save reports; global search across leads, accounts, opportunities, units, and bookings.")

# ── New capabilities built recently ──
h1("5. Configurable & Extensible Capabilities")

h2("5.1 Custom Fields — Object Manager")
bullet("Admin-defined fields", "Super Admins add fields to Account, Opportunity, Lead, and Site Visit at runtime — no database migration required (a metadata-driven design).")
bullet("12 data types", "Text, Text Area, Number, Currency, Percent, Checkbox, Date, Date/Time, Email, Phone, URL, and Picklist (with admin-defined options).")
bullet("Page-layout ordering", "Drag-and-drop (or arrow buttons) to control the display order of custom fields on forms and detail drawers.")
bullet("Everywhere it matters", "Custom fields render on the create forms and detail drawers of each object, with type validation enforced server-side.")

h2("5.2 Theming")
bullet("Six UI themes", "Indigo (default), Claude (warm clay/cream), Emerald, Ocean, Rose, and a true dark mode (Midnight) — switchable from the top bar and persisted per user.")
bullet("Token-driven", "Themes are CSS-variable palettes consumed by Tailwind, so a theme switch instantly re-skins the entire application.")

h2("5.3 Channel-Partner (Broker) Portal Integration")
bullet("Server-to-server intake API", "A secure /api/partner endpoint (API-key authenticated) accepts partner submissions without a user login.")
bullet("Lead & registration flow", "Broker-submitted leads become CRM Leads (source = BROKER, auto-assigned owner, broker credited in the timeline); registrations become CHANNEL_PARTNER accounts with a primary contact.")
bullet("Dual-write", "The broker portal writes to both its existing system and Reality CRM; the CRM write is authoritative for success while the legacy write is best-effort.")
bullet("Status mirroring", "Lead status changes in the portal mirror onto the correlated CRM lead's stage, recorded as a synced timeline note.")

# ── Architecture ──
h1("6. Architecture & Engineering Notes")
bullet("Backend", "Express + Prisma + PostgreSQL, layered as config / lib / middleware / modules; a permission cache is loaded at boot.")
bullet("Frontend", "React 18 + Vite + TypeScript, React Query for server state, Zustand for auth/session, Tailwind CSS for styling.")
bullet("Money handling", "All amounts are BigInt paise end-to-end; JSON emits them as strings and the client parses losslessly.")
bullet("Documents", "PDFs (allotment letter, quotation, demand, receipt) are rendered as branded print-ready HTML to keep headless browsers out of the API.")
bullet("Auditing & numbering", "Append-only audit on every mutation; atomic INSERT … ON CONFLICT counters for gap-free document numbers.")
bullet("Integrations", "Notifications (email/SMS/in-app), S3 signed-URL document storage, and the broker-portal partner API.")

doc.add_paragraph()
foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = foot.add_run("Reality CRM — Feature Guide · Generated for internal review")
r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = MUT

doc.save(OUT)
print("SAVED ->", OUT)
